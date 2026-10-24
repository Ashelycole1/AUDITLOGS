"""Organizations: isolation, feature switches, sign-up and platform administration
(auditlogs-additions_final.md 3 to 7, acceptance checklist 15)."""

import re

import pytest
from django.core import mail
from django.urls import reverse
from django.utils import timezone
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.auditlog.models import AuditLog
from apps.auditlog.services import record, verify_chain
from apps.organizations.models import Organization, OrganizationFeature

from .conftest import PASSWORD
from .test_permission_matrix import MATRIX, call

pytestmark = pytest.mark.django_db

SECRET = "ISOLATION-SECRET"


# Isolation ---------------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "method", "target", "allowed"),
    [m for m in MATRIX if m[2] is not None and m[3] - {"SYS"}],
    ids=[f"{m[1]} {m[0]}" for m in MATRIX if m[2] is not None and m[3] - {"SYS"}],
)
def test_another_organizations_record_is_not_found(
    name, method, target, allowed, client_for, make_user, make_org, url_for
):
    """Direct URLs and guessed ids of another organization's records answer 404, for every role."""
    url = url_for(name, target)  # made in the default organization
    other = make_org("Other organisation")
    for role in sorted(allowed - {"SYS"}):
        response = call(client_for(make_user(role, organization=other)), method, url)
        assert response.status_code == 404, (role, response.status_code, response.content[:200])


@pytest.fixture
def secret_data(make_user, make_query, make_report):
    """Recognisable records in the default organization that nobody elsewhere may see or count."""
    from apps.followup.models import Comment
    from apps.reports.models import Criterion

    ia = make_user("IA", email="isolation.person@secret.test")
    ia.full_name = f"{SECRET} Person"
    ia.save()
    report = make_report()
    report.title = f"{SECRET} report"
    report.save()
    criterion = Criterion.objects.create(report=report, title=f"{SECRET} criterion")
    query = make_query(
        assignee=ia,
        reference=f"{SECRET}-1",
        title=f"{SECRET} title",
        department=f"{SECRET} department",
        criterion=criterion,
        status="Assigned",
    )
    Comment.objects.create(query=query, author=make_user("CIA"), author_role="CIA", body=f"{SECRET} comment")
    record(
        actor=make_user("CIA"), action="test.secret", object_type="Query", object_id=query.pk, reason=SECRET
    )
    Organization.objects.filter(pk=query.organization_id).update(
        logo=SECRET.encode(), logo_updated_at=timezone.now()
    )
    return query


@pytest.mark.parametrize(
    ("name", "method", "allowed"),
    [(m[0], m[1], m[3]) for m in MATRIX if m[2] is None and m[1] == "get" and m[3] - {"SYS"}],
    ids=[m[0] for m in MATRIX if m[2] is None and m[1] == "get" and m[3] - {"SYS"}],
)
def test_lists_and_counts_hold_only_the_own_organization(
    name, method, allowed, secret_data, client_for, make_user, make_org
):
    other = make_org("Other organisation")
    # The data subject search is tried with the very term that would find the secret records.
    target = reverse(name) + (f"?q={SECRET}" if name == "subject-search" else "")
    for role in sorted(allowed - {"SYS"}):
        response = call(client_for(make_user(role, organization=other)), method, target)
        if name == "organization-logo":
            # The secret organization has a logo; this one has none, so it gets nothing at all.
            assert response.status_code == 404, role
            continue
        assert response.status_code == 200, (role, response.content[:200])
        body = response.content.decode("utf-8", errors="replace")
        # The search echoes the term it was given; nothing else may carry it.
        body = body.replace(f'"term":"{SECRET}"', "")
        assert SECRET not in body, (role, name)
        assert str(secret_data.pk) not in body, (role, name)


def test_widget_counts_are_per_organization(secret_data, client_for, make_user, make_org):
    other = make_org("Other organisation")
    mine = client_for(make_user("CIA", organization=secret_data.organization))
    theirs = client_for(make_user("CIA", organization=other))
    assert mine.get(reverse("dashboard-status-summary")).json()["total"] == 1
    assert theirs.get(reverse("dashboard-status-summary")).json()["total"] == 0
    assert theirs.get(reverse("dashboard-registered-numbers")).json()["queries"] == 0


def _export_text(response, fmt: str) -> str:
    """The readable text of an export: every cell of the workbook, or the PDF's uncompressed page text."""
    import io

    from openpyxl import load_workbook

    if fmt == "xlsx":
        book = load_workbook(io.BytesIO(response.content))
        return " ".join(
            str(c.value) for ws in book.worksheets for row in ws.iter_rows() for c in row if c.value
        )
    return response.content.decode("latin-1")


@pytest.mark.parametrize("fmt", ["xlsx", "pdf"])
def test_exports_hold_only_the_own_organization(
    fmt, secret_data, client_for, make_user, make_org, make_query, monkeypatch
):
    from reportlab import rl_config

    monkeypatch.setattr(rl_config, "pageCompression", 0)
    other = make_org("Other organisation")
    cia = client_for(make_user("CIA", organization=other))
    # A record of their own, to show the export would have shown a match.
    from apps.imports.models import ImportBatch
    from apps.queries.models import Query

    batch = ImportBatch.objects.create(
        imported_by=cia.user,
        source_filename="x.csv",
        source_file_path="unused",
        source_sha256="0" * 64,
        status="Confirmed",
    )
    Query.objects.create(
        reference="OWN-1", batch=batch, title="OWNMARKER title", description_original="d", description_current="d",
        risk_rating_original="Low", risk_rating_current="Low", deadline="2030-01-01", responsible_officer="o", department="d",
    )  # fmt: skip
    response = cia.get(reverse("queries-export"), {"export_format": fmt})
    assert response.status_code == 200
    text = _export_text(response, fmt)
    assert "OWNMARKER" in text
    assert SECRET not in text


def test_cannot_assign_to_another_organizations_auditor(client_for, make_user, make_org, make_query):
    query = make_query()
    stranger = make_user("IA", organization=make_org("Other organisation"))
    response = client_for("CIA").post(
        reverse("queries-assign", kwargs={"pk": query.pk}), {"assignee": str(stranger.pk)}, format="json"
    )
    assert response.status_code == 400
    query.refresh_from_db()
    assert query.current_assignee_id is None


def test_records_cannot_link_two_organizations(make_user, make_org, make_query):
    from apps.organizations.models import CrossOrganizationError
    from apps.reports.models import AuditReport, Criterion

    query = make_query()
    other = make_org("Other organisation")
    foreign = AuditReport.objects.create(
        organization=other, title="Theirs", created_by=make_user("CIA", organization=other)
    )
    criterion = Criterion.objects.create(report=foreign, title="Their criterion")
    query.criterion = criterion
    with pytest.raises(CrossOrganizationError):
        query.save()


def test_same_reference_may_exist_in_two_organizations(make_query, make_user, make_org, make_batch):
    from apps.imports.models import ImportBatch
    from apps.queries.models import Query

    make_query(reference="IA/2026/001")
    other = make_org("Other organisation")
    batch = ImportBatch.objects.create(
        imported_by=make_user("CIA", organization=other),
        source_filename="x.csv",
        source_file_path="unused",
        source_sha256="0" * 64,
        status="Confirmed",
    )
    Query.objects.create(
        reference="IA/2026/001",
        batch=batch,
        title="t",
        description_original="d",
        description_current="d",
        risk_rating_original="Low",
        risk_rating_current="Low",
        deadline="2030-01-01",
        responsible_officer="o",
        department="d",
    )
    assert Query.objects.filter(reference="IA/2026/001").count() == 2


def test_audit_log_and_verification_speak_only_of_the_own_organization(client_for, make_user, make_org):
    other = make_org("Other organisation")
    record(actor=make_user("CIA", organization=other), action="test.theirs", object_type="T", object_id="1")
    cia = make_user("CIA")
    record(actor=cia, action="test.mine", object_type="T", object_id="1")
    client = client_for(cia)
    actions = [e["action"] for e in client.get(reverse("audit-log")).json()["results"]]
    assert "test.mine" in actions and "test.theirs" not in actions
    result = client.post(reverse("audit-log-verify")).json()
    assert result["ok"] is True
    mine = AuditLog.objects.filter(organization=cia.organization).count()
    assert result["checked"] == mine - 1  # the verify entry itself is written after the walk


def test_the_organization_is_sealed_into_each_entry(make_user, make_org):
    from django.db import connection

    from .test_auditlog import drop_triggers

    entry = record(actor=make_user("CIA"), action="test.event", object_type="T", object_id="1")
    other = make_org("Other organisation")
    drop_triggers()
    with connection.cursor() as cursor:
        cursor.execute(
            "UPDATE auditlog_auditlog SET organization_id = %s WHERE seq = %s", [other.pk.hex, entry.seq]
        )
    assert verify_chain().broken_at == entry.seq


# Feature switches -------------------------------------------------------------------------------------


def switch(org, key: str, enabled: bool) -> None:
    OrganizationFeature.objects.update_or_create(
        organization=org, feature_key=key, defaults={"enabled": enabled}
    )


@pytest.mark.parametrize(
    ("key", "url"),
    [
        ("component.status_pie", "dashboard-status-summary"),
        ("component.registered_numbers", "dashboard-registered-numbers"),
        ("component.system_stats", "dashboard-system-stats"),
    ],
)
def test_a_switched_off_widget_is_refused_by_the_server(key, url, org, client_for):
    client = client_for("CIA")
    assert client.get(reverse(url)).status_code == 200
    switch(org, key, False)
    assert client.get(reverse(url)).status_code == 403
    assert key not in client.get(reverse("auth-me")).json()["organization"]["features"]


def test_a_switched_off_dashboard_shuts_out_its_role(org, client_for, make_user):
    ceo = make_user("CEO")
    client = client_for(ceo)
    assert client.get(reverse("dashboard-queries")).status_code == 200
    ceo.is_active = False
    ceo.save()
    switch(org, "dashboard.CEO", False)
    ceo.is_active = True
    ceo.save()
    assert client.get(reverse("dashboard-queries")).status_code == 403
    assert client.get(reverse("auth-me")).status_code == 403
    login = APIClient().post(reverse("auth-login"), {"email": ceo.email, "password": PASSWORD}, format="json")
    assert login.status_code == 403
    assert login.json()["code"] == "role"


def test_ict_cannot_give_a_role_whose_dashboard_is_off(org, client_for, make_user):
    switch(org, "dashboard.SIA", False)
    ict = client_for("ICT")
    response = ict.post(
        reverse("users-list"),
        {"email": "new@auditlogs.test", "full_name": "New", "role": "SIA"},
        format="json",
    )
    assert response.status_code == 400
    assert "not enabled" in str(response.json())
    ia = make_user("IA")
    response = ict.patch(
        reverse("users-detail", kwargs={"pk": ia.pk}),
        {"role": "SIA", "reason": "Promoted after review"},
        format="json",
    )
    assert response.status_code == 400
    assert (
        ict.post(
            reverse("users-list"),
            {"email": "n2@auditlogs.test", "full_name": "N", "role": "SYS"},
            format="json",
        ).status_code
        == 400
    )


def test_ict_manages_only_its_own_organization(client_for, make_user, make_org):
    other = make_org("Other organisation")
    stranger = make_user("IA", organization=other)
    ict = client_for("ICT")
    ids = {u["id"] for u in ict.get(reverse("users-list")).json()}
    assert str(stranger.pk) not in ids
    response = ict.post(
        reverse("users-list"), {"email": "x@auditlogs.test", "full_name": "X", "role": "IA"}, format="json"
    )
    assert User.objects.get(pk=response.json()["id"]).organization_id == ict.user.organization_id


@pytest.fixture
def sysadmin(client_for, make_user):
    return client_for(make_user("SYS"))


def test_features_safeguards_and_audit_trail(sysadmin, org, make_user):
    url = reverse("platform-organizations-features", kwargs={"pk": org.pk})
    assert sysadmin.put(url, {"features": {"dashboard.IA": False}}, format="json").status_code == 409
    make_user("SIA")
    response = sysadmin.put(url, {"features": {"dashboard.SIA": False}}, format="json")
    assert response.status_code == 409
    assert "Senior Internal Auditor" in response.json()["detail"]

    response = sysadmin.put(
        url, {"preset": "core", "reason": "Client asked for the core package"}, format="json"
    )
    assert response.status_code == 409  # the SIA still blocks it
    User.objects.filter(organization=org, role="SIA").update(is_active=False)
    response = sysadmin.put(
        url, {"preset": "core", "reason": "Client asked for the core package"}, format="json"
    )
    assert response.status_code == 200
    assert response.json()["features"]["dashboard.SIA"] is False
    entry = AuditLog.objects.filter(action="organization.features").latest("seq")
    assert entry.organization_id == org.pk
    assert entry.old_value == {"dashboard.SIA": True} and entry.new_value == {"dashboard.SIA": False}
    assert entry.reason == "Client asked for the core package"
    # Switching it back on restores the role; nothing was deleted.
    assert sysadmin.put(url, {"features": {"dashboard.SIA": True}}, format="json").json()["features"][
        "dashboard.SIA"
    ]


def test_sysadmin_sees_counts_not_audit_content(sysadmin, secret_data):
    stats = sysadmin.get(reverse("platform-stats")).json()
    assert stats["queries"] == 1
    assert SECRET not in str(stats)
    for name in ("queries-list", "audit-log", "dashboard-status-summary", "reports-list"):
        assert sysadmin.get(reverse(name)).status_code == 403


# Organization status and sign-in ----------------------------------------------------------------------


def login(email: str, portal: str = "organization"):
    return APIClient().post(
        reverse("auth-login"), {"email": email, "password": PASSWORD, "portal": portal}, format="json"
    )


@pytest.mark.parametrize("org_status", ["Pending", "Suspended", "Deactivated", "Rejected"])
def test_people_of_a_closed_organization_cannot_sign_in(org_status, make_org, make_user, settings):
    settings.MFA_REQUIRED = False
    user = make_user("CIA", organization=make_org("Closed", status=org_status))
    response = login(user.email)
    assert response.status_code == 403
    assert response.json()["code"] == "organization"


def test_suspension_ends_open_sessions(org, client_for):
    client = client_for("CIA")
    assert client.get(reverse("auth-me")).status_code == 200
    org.status = "Suspended"
    org.save()
    assert client.get(reverse("auth-me")).status_code == 403


def test_sysadmin_signs_in_on_the_platform_page_only(make_user, settings):
    settings.MFA_REQUIRED = False
    admin = make_user("SYS")
    assert login(admin.email).json()["code"] == "portal"
    assert login(admin.email, "platform").status_code == 200
    cia = make_user("CIA")
    assert login(cia.email, "platform").json()["code"] == "portal"
    assert login(cia.email).status_code == 200


def test_session_must_belong_to_the_accounts_organization(client_for, make_user):
    client = client_for("CIA")
    session = client.session
    session["organization"] = "platform"
    session.save()
    assert client.get(reverse("auth-me")).status_code == 403


# Sign-up and approval ---------------------------------------------------------------------------------

SIGNUP = {
    "name": "Mbale Water Board",
    "country": "Uganda",
    "contact_person": "Agnes Auma",
    "email": "agnes.auma@mbalewater.test",
    "phone": "+256 700 000000",
}


def test_sign_up_waits_for_approval_then_gets_its_package(
    sysadmin, settings, django_capture_on_commit_callbacks
):
    settings.MFA_REQUIRED = False
    with django_capture_on_commit_callbacks(execute=True):
        response = APIClient().post(reverse("signup"), {**SIGNUP, "password": PASSWORD}, format="json")
    assert response.status_code == 201, response.content
    org = Organization.objects.get(name=SIGNUP["name"])
    assert org.status == "Pending"
    admin = User.objects.get(email=SIGNUP["email"])
    assert (admin.role, admin.organization_id) == ("ICT", org.pk)
    assert login(admin.email).json()["code"] == "organization"

    token = re.search(r"token=([^&\s]+)", response.json()["verify_link"]).group(1)
    assert APIClient().post(reverse("signup-verify"), {"token": token}, format="json").status_code == 200
    org.refresh_from_db()
    assert org.email_verified_at is not None

    approve = reverse("platform-organizations-approve", kwargs={"pk": org.pk})
    with django_capture_on_commit_callbacks(execute=True):
        response = sysadmin.post(approve, {"preset": "core"}, format="json")
    assert response.status_code == 200
    assert response.json()["status"] == "Active"
    assert response.json()["features"]["dashboard.SIA"] is False
    assert login(admin.email).status_code == 200
    # A confirmation link on sign-up and the approval notice, both to the organization's address.
    subjects = [m.subject for m in mail.outbox if SIGNUP["email"] in m.to]
    assert any("confirm your email" in s for s in subjects) and any("approved" in s for s in subjects)


def test_sign_up_refuses_a_taken_email_and_a_weak_password(make_user):
    taken = make_user("IA")
    response = APIClient().post(
        reverse("signup"), {**SIGNUP, "email": taken.email, "password": PASSWORD}, format="json"
    )
    assert response.status_code == 400
    response = APIClient().post(reverse("signup"), {**SIGNUP, "password": "short"}, format="json")
    assert response.status_code == 400
    assert not Organization.objects.filter(name=SIGNUP["name"]).exists()


def test_lifecycle_needs_reasons_and_follows_the_allowed_moves(sysadmin, make_org):
    org = make_org("Lifecycle", status="Active")
    suspend = reverse("platform-organizations-suspend", kwargs={"pk": org.pk})
    assert sysadmin.post(suspend, {}, format="json").status_code == 409
    assert (
        sysadmin.post(suspend, {"reason": "Unpaid subscription invoice"}, format="json").json()["status"]
        == "Suspended"
    )
    approve = reverse("platform-organizations-approve", kwargs={"pk": org.pk})
    assert sysadmin.post(approve, {}, format="json").status_code == 409
    reactivate = reverse("platform-organizations-reactivate", kwargs={"pk": org.pk})
    assert (
        sysadmin.post(reactivate, {"reason": "Invoice paid in full"}, format="json").json()["status"]
        == "Active"
    )
    actions = list(AuditLog.objects.filter(organization=org).values_list("action", flat=True))
    assert actions == ["organization.suspend", "organization.reactivate"]


def test_sysadmin_creates_an_organization_and_its_first_admin(sysadmin, settings):
    settings.MFA_REQUIRED = False
    response = sysadmin.post(
        reverse("platform-organizations-list"),
        {**SIGNUP, "preset": "full", "admin_email": "it@mbalewater.test", "admin_name": "Ivan Ochieng"},
        format="json",
    )
    assert response.status_code == 201, response.content
    body = response.json()
    assert body["status"] == "Active" and body["administrators"][0]["email"] == "it@mbalewater.test"
    link = body["admin"]["set_password_link"]
    uid, token = re.search(r"uid=([^&]+)&token=([^&\s]+)", link).groups()
    set_password = reverse("auth-set-password")
    assert (
        APIClient()
        .post(set_password, {"uid": uid, "token": token, "password": PASSWORD}, format="json")
        .status_code
        == 200
    )
    # The link works once.
    assert (
        APIClient()
        .post(set_password, {"uid": uid, "token": token, "password": PASSWORD}, format="json")
        .status_code
        == 400
    )
    assert login("it@mbalewater.test").status_code == 200

    # A reset stops the old password and sends a fresh link.
    org_id = body["id"]
    reset = sysadmin.post(
        reverse("platform-organizations-admin", kwargs={"pk": org_id}),
        {"email": "it@mbalewater.test", "full_name": "Ivan Ochieng", "reason": "Administrator lost access"},
        format="json",
    )
    assert reset.status_code == 200
    assert login("it@mbalewater.test").status_code == 400


def test_admin_email_of_another_account_is_refused(sysadmin, make_user, make_org):
    taken = make_user("CIA")
    org = make_org("Elsewhere")
    response = sysadmin.post(
        reverse("platform-organizations-admin", kwargs={"pk": org.pk}),
        {"email": taken.email, "full_name": "Someone"},
        format="json",
    )
    assert response.status_code == 409
    taken.refresh_from_db()
    assert taken.organization_id != org.pk


# Organization profile: name, details and logo (owner, 2026-10-25) --------------------------------------


def picture(fmt: str = "JPEG", size=(1200, 400)) -> bytes:
    import io

    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", size, (20, 90, 160)).save(out, format=fmt)
    return out.getvalue()


def upload_logo(
    client, org, content: bytes, name: str = "logo.jpg", reason: str = "New brand from the client"
):
    from django.core.files.uploadedfile import SimpleUploadedFile

    return client.post(
        reverse("platform-organizations-logo", kwargs={"pk": org.pk}),
        {"file": SimpleUploadedFile(name, content), "reason": reason},
        format="multipart",
    )


def test_sysadmin_changes_the_details_the_organization_sees(sysadmin, org, client_for):
    response = sysadmin.patch(
        reverse("platform-organizations-detail", kwargs={"pk": org.pk}),
        {
            "name": "Mbale Water Board",
            "address": "Plot 4, Republic Street, Mbale",
            "website": "https://mbalewater.test",
            "reason": "Name and address changed by the client",
        },
        format="json",
    )
    assert response.status_code == 200, response.content
    mine = client_for("PIA").get(reverse("auth-me")).json()["organization"]
    assert (mine["name"], mine["address"], mine["website"]) == (
        "Mbale Water Board",
        "Plot 4, Republic Street, Mbale",
        "https://mbalewater.test",
    )
    entry = AuditLog.objects.filter(action="organization.update").latest("seq")
    assert entry.new_value["address"] == "Plot 4, Republic Street, Mbale" and entry.organization_id == org.pk


def test_logo_is_redrawn_as_a_small_png_and_shown_to_the_organization(sysadmin, org, client_for, make_org):
    import io

    from PIL import Image

    response = upload_logo(sysadmin, org, picture())
    assert response.status_code == 200, response.content
    assert response.json()["logo_url"].startswith(
        reverse("platform-organizations-logo", kwargs={"pk": org.pk}) + "?v="
    )

    ia = client_for("IA")
    url = ia.get(reverse("auth-me")).json()["organization"]["logo_url"]
    assert url.startswith(reverse("organization-logo") + "?v=")
    served = ia.get(reverse("organization-logo"))
    assert served.status_code == 200 and served["Content-Type"] == "image/png"
    assert served["X-Content-Type-Options"] == "nosniff"
    with Image.open(io.BytesIO(served.content)) as img:
        assert img.format == "PNG" and max(img.size) == 512

    entry = AuditLog.objects.filter(action="organization.logo").latest("seq")
    assert entry.organization_id == org.pk and len(entry.new_value["logo_sha256"]) == 64
    assert entry.reason == "New brand from the client"

    # Another organization has its own (no) logo, never this one.
    other = client_for(make_user_in(make_org("Other organisation")))
    assert other.get(reverse("organization-logo")).status_code == 404
    assert other.get(reverse("auth-me")).json()["organization"]["logo_url"] is None


def make_user_in(org):
    return User.objects.create_user(
        email=f"ia.{org.pk.hex[:8]}@org.test",
        full_name="Other IA",
        role="IA",
        password=PASSWORD,
        organization=org,
    )


@pytest.mark.parametrize(
    ("name", "content"),
    [
        ("logo.svg", b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'),
        ("logo.png", b"not a picture"),
        ("logo.bmp", picture("BMP", (10, 10))),
    ],
)
def test_logo_must_be_a_plain_picture(sysadmin, org, name, content):
    assert upload_logo(sysadmin, org, content, name).status_code == 400
    assert Organization.objects.get(pk=org.pk).logo is None


def test_logo_change_needs_a_reason_and_can_be_taken_down(sysadmin, org, client_for):
    assert upload_logo(sysadmin, org, picture(), reason="").status_code == 400
    assert upload_logo(sysadmin, org, picture("PNG", (64, 64))).status_code == 200
    url = reverse("platform-organizations-logo", kwargs={"pk": org.pk})
    assert sysadmin.delete(url, {}, format="json").status_code == 409
    response = sysadmin.delete(url, {"reason": "Old logo withdrawn"}, format="json")
    assert response.status_code == 200 and response.json()["logo_url"] is None
    assert client_for("CEO").get(reverse("organization-logo")).status_code == 404
    assert AuditLog.objects.filter(action="organization.logo_remove").exists()
    assert verify_chain().ok
