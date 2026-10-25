"""Query record, audit reports, Word import, PDF export, original files and the status pie chart
(auditlogs-additions_final.md 8 to 13)."""

import io
from datetime import date, timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.auditlog.models import AuditLog
from apps.imports.parsing import split_action_party_time
from apps.queries.models import Query, QueryRevision
from apps.reports.models import AuditReport, Criterion

pytestmark = pytest.mark.django_db

FUTURE = date.today() + timedelta(days=60)
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def word_report(*, extra_query: bool = True) -> bytes:
    """A Word audit report as the template describes it: a header table, criteria as headings, and one
    two-column table per query."""
    import docx

    document = docx.Document()
    document.add_heading("Internal audit of procurement, 2026", level=0)
    header = document.add_table(rows=0, cols=2)
    for label, value in (
        ("Report title", "Procurement audit 2026"),
        ("Audit period", "July 2025 to June 2026"),
        ("Report date", "15/09/2026"),
        ("Prepared by", "Internal Audit Unit"),
    ):
        row = header.add_row().cells
        row[0].text, row[1].text = label, value

    def query_table(fields):
        table = document.add_table(rows=0, cols=2)
        for label, value in fields:
            row = table.add_row().cells
            row[0].text, row[1].text = label, value

    document.add_heading("Criterion 1: Public Procurement and Disposal Act, section 58", level=1)
    document.add_paragraph("Contracts above the threshold must go through open bidding.")
    query_table(
        [
            ("Reference", "PR/2026/001"),
            ("Title", "Contracts split to avoid open bidding"),
            ("Query", "Three contracts for the same works were split below the threshold."),
            ("Management Response", "Management agrees and will consolidate future works."),
            ("Risk Rating", "High"),
            ("Recommendation", "Plan works annually and bid them as one lot."),
            ("Audit Observation", "Invoices show the same supplier on consecutive days."),
            ("Policy Title", "Procurement Policy 2024"),
            ("Policy Reference", "Clause 4.2: no splitting of requirements."),
            ("Action Party & Time", f"Head of Procurement, {FUTURE:%d/%m/%Y}"),
            ("Department", "Procurement"),
            ("Notable Numbers", "UGX 148,500,000 across 3 contracts"),
        ]
    )
    if extra_query:
        document.add_heading("Criterion 2: Treasury Instructions 2017", level=1)
        query_table(
            [
                ("Reference", "PR/2026/002"),
                ("Query", "Advances were not accounted for within 60 days. Some remain open after a year."),
                ("Risk Rating", "Medium"),
                ("Action Party", "Finance Manager"),
                ("Due date", FUTURE.isoformat()),
                ("Department", "Finance"),
            ]
        )
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def upload(client, content: bytes, name: str = "audit-report.docx", **extra):
    return client.post(
        reverse("imports-list"),
        {"file": SimpleUploadedFile(name, content, content_type=DOCX), **extra},
        format="multipart",
    )


def test_action_party_and_time_split():
    assert split_action_party_time("Finance Manager, 31/01/2027") == ("Finance Manager", "31/01/2027")
    assert split_action_party_time("Head of ICT by 2027-03-01") == ("Head of ICT", "2027-03-01")
    assert split_action_party_time("Stores Supervisor") == ("Stores Supervisor", "")


def test_word_report_imports_the_full_record_with_its_report_and_criteria(client_for):
    cia = client_for("CIA")
    response = upload(cia, word_report())
    assert response.status_code == 201, response.content
    preview = response.json()
    assert preview["error_count"] == 0, preview["rows"]
    assert preview["report_meta"]["title"] == "Procurement audit 2026"
    second = preview["rows"][1]
    assert second["data"]["title"] == "Advances were not accounted for within 60 days."
    assert "first sentence" in second["warnings"][0]

    assert cia.post(reverse("imports-confirm", kwargs={"pk": preview["id"]})).status_code == 200
    report = AuditReport.objects.get()
    assert (report.title, report.audit_period, report.report_date) == (
        "Procurement audit 2026",
        "July 2025 to June 2026",
        date(2026, 9, 15),
    )
    assert list(report.criteria.values_list("title", flat=True)) == [
        "Public Procurement and Disposal Act, section 58",
        "Treasury Instructions 2017",
    ]
    assert (
        report.criteria.first().description == "Contracts above the threshold must go through open bidding."
    )
    q = Query.objects.get(reference="PR/2026/001")
    assert q.criterion.title == "Public Procurement and Disposal Act, section 58"
    assert q.management_response.startswith("Management agrees")
    assert q.recommendation.startswith("Plan works annually")
    assert q.audit_observation.startswith("Invoices show")
    assert (q.policy_title, q.policy_reference) == (
        "Procurement Policy 2024",
        "Clause 4.2: no splitting of requirements.",
    )
    assert (q.responsible_officer, q.deadline) == ("Head of Procurement", FUTURE)
    assert q.notable_numbers == "UGX 148,500,000 across 3 contracts"
    assert q.risk_rating_original == "High"
    assert AuditLog.objects.filter(action="report.create").count() == 1
    assert AuditLog.objects.filter(action="criterion.create").count() == 2


def test_word_import_into_a_chosen_report(client_for):
    cia = client_for("CIA")
    report = AuditReport.objects.create(
        organization=cia.user.organization, title="Chosen", created_by=cia.user
    )
    batch = upload(cia, word_report(extra_query=False), report=str(report.pk)).json()
    cia.post(reverse("imports-confirm", kwargs={"pk": batch["id"]}))
    assert AuditReport.objects.count() == 1
    assert Query.objects.get().criterion.report_id == report.pk


def test_a_word_file_without_queries_is_refused(client_for):
    import docx

    document = docx.Document()
    document.add_paragraph("Nothing here.")
    out = io.BytesIO()
    document.save(out)
    response = upload(client_for("CIA"), out.getvalue())
    assert response.status_code == 400
    assert "No queries were found" in response.json()["detail"]
    assert upload(client_for("CIA"), b"not a zip").status_code == 400
    assert upload(client_for("CIA"), b"old", name="old.doc").status_code == 400


def test_csv_with_the_record_columns_and_criteria(client_for):
    header = (
        "reference,title,description,risk_rating,action party & time,department,"
        "recommendation,audit observation,policy title,policy reference,notable numbers,management response,criterion"
    )
    line = (
        f'IA/2026/500,Fuel cards shared,Cards used by several drivers.,Low,"Transport Officer, {FUTURE:%d/%m/%Y}",'
        "Transport,Issue one card per driver,Card log shows 4 users,Fleet Policy,Section 3,12 cards,Agreed,Fleet rules"
    )
    content = "\n".join([header, line]).encode()
    cia = client_for("CIA")
    preview = cia.post(
        reverse("imports-list"), {"file": SimpleUploadedFile("q.csv", content)}, format="multipart"
    ).json()
    assert preview["error_count"] == 0, preview["rows"]
    cia.post(reverse("imports-confirm", kwargs={"pk": preview["id"]}))
    q = Query.objects.get()
    assert (q.recommendation, q.notable_numbers, q.management_response) == (
        "Issue one card per driver",
        "12 cards",
        "Agreed",
    )
    assert q.criterion.title == "Fleet rules"
    assert q.criterion.report.title == "Audit report: q.csv"


def test_cia_edits_the_record_fields_with_revisions(client_for, make_query):
    cia = client_for("CIA")
    query = make_query()
    report = AuditReport.objects.create(organization=query.organization, title="R", created_by=cia.user)
    criterion = Criterion.objects.create(report=report, title="Finance Act")
    url = reverse("queries-detail", kwargs={"pk": query.pk})
    response = cia.patch(
        url,
        {
            "recommendation": "Count stock monthly",
            "criterion": str(criterion.pk),
            "reason": "Agreed at the exit meeting",
        },
        format="json",
    )
    assert response.status_code == 200, response.content
    assert response.json()["criterion_title"] == "Finance Act"
    assert response.json()["recommendation"] == "Count stock monthly"
    revisions = {r.field: r.new_value for r in QueryRevision.objects.filter(query=query)}
    assert revisions == {"recommendation": "Count stock monthly", "criterion": "Finance Act"}


def test_a_criterion_of_another_organization_cannot_be_set(client_for, make_query, make_org, make_user):
    query = make_query()
    other = make_org("Other organisation")
    foreign = AuditReport.objects.create(
        organization=other, title="R", created_by=make_user("CIA", organization=other)
    )
    criterion = Criterion.objects.create(report=foreign, title="Theirs")
    response = client_for("CIA").patch(
        reverse("queries-detail", kwargs={"pk": query.pk}),
        {"criterion": str(criterion.pk), "reason": "Trying another organization"},
        format="json",
    )
    assert response.status_code == 400
    query.refresh_from_db()
    assert query.criterion_id is None


def test_ceo_cannot_change_anything(client_for, make_query, make_user):
    ceo = client_for("CEO")
    query = make_query()
    url = reverse("queries-detail", kwargs={"pk": query.pk})
    assert ceo.get(url).status_code == 200
    assert (
        ceo.patch(url, {"title": "Changed", "reason": "CEO trying to edit"}, format="json").status_code == 403
    )
    assert (
        ceo.post(
            reverse("queries-assign", kwargs={"pk": query.pk}),
            {"assignee": str(make_user("IA").pk)},
            format="json",
        ).status_code
        == 403
    )
    assert (
        ceo.post(
            reverse("queries-change-status", kwargs={"pk": query.pk}),
            {"to_status": "Closed", "reason": "x" * 12},
            format="json",
        ).status_code
        == 403
    )
    assert ceo.post(reverse("reports-list"), {"title": "CEO report"}, format="json").status_code == 403
    query.refresh_from_db()
    assert query.title != "Changed"


def test_reports_api_and_pdf(client_for, make_query):
    cia = client_for("CIA")
    created = cia.post(
        reverse("reports-list"), {"title": "Stores audit", "audit_period": "FY 2025/26"}, format="json"
    )
    assert created.status_code == 201
    report_id = created.json()["id"]
    criterion = cia.post(
        reverse("reports-criteria", kwargs={"pk": report_id}),
        {"title": "Stores Manual", "description": "Counts"},
        format="json",
    )
    assert criterion.status_code == 201
    assert (
        cia.post(
            reverse("reports-criteria", kwargs={"pk": report_id}), {"title": "Stores Manual"}, format="json"
        ).status_code
        == 409
    )
    query = make_query()
    cia.patch(
        reverse("queries-detail", kwargs={"pk": query.pk}),
        {"criterion": criterion.json()["id"], "reason": "Grouped under the manual"},
        format="json",
    )

    detail = cia.get(reverse("reports-detail", kwargs={"pk": report_id})).json()
    assert detail["criteria"][0]["items"][0]["reference"] == query.reference
    assert detail["summary"]["total"] == 1

    pdf = cia.get(reverse("reports-pdf", kwargs={"pk": report_id}))
    assert pdf.status_code == 200
    assert pdf["Content-Type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")
    assert AuditLog.objects.filter(action="export.report").exists()

    rename = cia.patch(
        reverse("reports-detail", kwargs={"pk": report_id}), {"title": "Stores audit 2026"}, format="json"
    )
    assert rename.json()["title"] == "Stores audit 2026"
    edit = cia.patch(
        reverse("reports-update-criterion", kwargs={"pk": report_id, "criterion_id": criterion.json()["id"]}),
        {"title": "Stores Manual 2024"},
        format="json",
    )
    assert edit.json()["title"] == "Stores Manual 2024"


def test_pdf_export_carries_the_record_and_the_organization(client_for, make_query, monkeypatch):
    from reportlab import rl_config

    monkeypatch.setattr(rl_config, "pageCompression", 0)
    cia = client_for("CIA")
    make_query(recommendation="Reconcile the float weekly", notable_numbers="UGX 2,000,000")
    response = cia.get(reverse("queries-export"), {"export_format": "pdf"})
    assert response.status_code == 200
    text = response.content.decode("latin-1")
    for expected in (
        "Reconcile the float weekly",
        "UGX 2,000,000",
        "Test organisation",
        "Queries worked on",
        "8. Notable numbers",
    ):
        assert expected in text, expected
    assert cia.get(reverse("queries-export"), {"export_format": "docx"}).status_code == 409


def test_pia_pdf_export_waits_for_approval(client_for, make_user):
    pia = client_for("PIA")
    created = pia.post(
        reverse("export-requests-list"),
        {"filters": {}, "format": "pdf", "reason": "Board pack for October"},
        format="json",
    )
    assert created.status_code == 201
    export_id = created.json()["id"]
    assert pia.get(reverse("export-requests-download", kwargs={"pk": export_id})).status_code == 409
    client_for("CIA").post(reverse("export-requests-approve", kwargs={"pk": export_id}), {}, format="json")
    response = pia.get(reverse("export-requests-download", kwargs={"pk": export_id}))
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_original_file_downloads_unchanged_and_is_logged(client_for, file_storage):
    cia = client_for("CIA")
    content = word_report()
    batch_id = upload(cia, content).json()["id"]
    for role in ("CIA", "PIA", "CEO"):
        response = client_for(role).get(reverse("imports-original", kwargs={"pk": batch_id}))
        assert response.status_code == 200, role
        assert response.content == content
        assert "audit-report.docx" in response["Content-Disposition"]
    assert AuditLog.objects.filter(action="import.download_original").count() == 3
    # Stored under the organization's own folder.
    from apps.imports.models import ImportBatch

    batch = ImportBatch.objects.get(pk=batch_id)
    assert str(batch.organization_id) in batch.source_file_path


def test_a_changed_original_is_refused(client_for, file_storage):
    from pathlib import Path

    from cryptography.fernet import Fernet
    from django.conf import settings

    from apps.imports.models import ImportBatch

    cia = client_for("CIA")
    batch_id = upload(cia, word_report()).json()["id"]
    batch = ImportBatch.objects.get(pk=batch_id)
    Path(batch.source_file_path).write_bytes(Fernet(settings.FILE_ENCRYPTION_KEY).encrypt(b"swapped"))
    assert cia.get(reverse("imports-original", kwargs={"pk": batch_id})).status_code == 409


def test_status_pie_adds_up(client_for, make_query, make_user):
    ia = make_user("IA")
    make_query()  # not allocated
    make_query(assignee=ia, status="Assigned")  # allocated, not worked on
    make_query(assignee=ia, status="In Progress")  # worked on
    make_query(assignee=ia, status="Closed")  # worked on
    body = client_for("IA").get(reverse("dashboard-status-summary")).json()
    assert (body["total"], body["worked_on"], body["not_worked_on"], body["allocated"]) == (4, 2, 2, 3)
    assert [s["count"] for s in body["slices"]] == [2, 1, 1]
    assert sum(s["count"] for s in body["slices"]) == body["total"]
    assert body["percent"]["worked_on"] == 50.0
    # A slice filters the query list.
    cia = client_for("CIA")
    assert cia.get(reverse("queries-list"), {"worked": "yes"}).json()["count"] == 2
    assert cia.get(reverse("queries-list"), {"worked": "no", "assignee": "any"}).json()["count"] == 1


def test_status_pie_counts_queries_per_person_without_their_details(client_for, make_query, make_user):
    ia, sia = make_user("IA"), make_user("SIA")
    make_query()
    make_query(assignee=ia, status="Assigned", title="Secret payroll finding")
    make_query(assignee=ia, status="In Progress")
    make_query(assignee=sia, status="Closed")
    # Every role sees who holds how many, the IA included, even for queries it cannot open.
    body = client_for("IA").get(reverse("dashboard-status-summary")).json()
    assert body["allocated_to"] == [
        {"id": str(ia.pk), "full_name": ia.full_name, "role": "IA", "count": 2, "worked_on": 1},
        {"id": str(sia.pk), "full_name": sia.full_name, "role": "SIA", "count": 1, "worked_on": 1},
    ]
    assert sum(p["count"] for p in body["allocated_to"]) == body["allocated"]
    # Counts only: nothing of a query itself.
    assert "Secret" not in str(body) and "reference" not in str(body)


def test_assignee_name_and_history(client_for, make_query, make_user):
    from apps.queries.models import Assignment

    first, second = make_user("IA"), make_user("SIA")
    query = make_query()
    cia = client_for("CIA")
    cia.post(reverse("queries-assign", kwargs={"pk": query.pk}), {"assignee": str(first.pk)}, format="json")
    cia.post(
        reverse("queries-assign", kwargs={"pk": query.pk}),
        {"assignee": str(second.pk), "reason": "Workload balancing"},
        format="json",
    )
    row = cia.get(reverse("queries-list")).json()["results"][0]
    assert row["assignee_name"] == second.full_name
    history = list(
        Assignment.objects.filter(query=query).order_by("assigned_at").values_list("assignee", "active")
    )
    assert history == [(first.pk, False), (second.pk, True)]
