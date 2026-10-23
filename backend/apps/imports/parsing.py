"""Reading and checking a final-queries file (design.md 4 and 13; auditlogs-additions_final.md 8.2 and 13).

Excel and CSV files are one table with a header row. A Word (.docx) audit report may carry a report header
table, criteria as headings ("Criterion 1: ..."), and its queries either as one two-column table per query
(field name, value) or as one table with a header row. Headers and field names are matched through
aliases; settings.IMPORT_HEADER_ALIASES adds more per field until the client's template is final.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from django.conf import settings
from openpyxl import load_workbook

FIELDS = ("reference", "title", "description", "risk_rating", "deadline", "responsible_officer", "department")
# The rest of the query record (section 8.2) and the criterion it sits under. Optional in every format.
OPTIONAL_FIELDS = (
    "management_response",
    "recommendation",
    "audit_observation",
    "policy_title",
    "policy_reference",
    "notable_numbers",
    "criterion",
)
# One cell holding both "who must act" and "by when", e.g. "Finance Manager, 31/01/2027".
ACTION_PARTY_TIME = "action_party_time"
# "Extreme" is the reference's other name for Critical (Standard Risk Matrix Color Mapping Reference).
RISKS = {"critical": "Critical", "extreme": "Critical", "high": "High", "medium": "Medium", "low": "Low"}
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %B %Y", "%d %b %Y")  # Ugandan files: day first.
MAX_LENGTH = {
    "reference": 50,
    "title": 300,
    "responsible_officer": 150,
    "department": 150,
    "policy_title": 300,
    "criterion": 300,
}
BUILTIN_ALIASES = {
    "reference": ["ref", "ref no", "query ref", "query reference", "reference number"],
    "title": ["query title"],
    "description": ["query", "query text", "query description"],
    "risk_rating": ["risk", "risk level"],
    "deadline": ["action date", "action time", "due date", "time"],
    "responsible_officer": ["action party", "responsible person"],
    "department": ["unit", "directorate"],
    "management_response": ["response", "management comment"],
    "recommendation": ["recommendations", "auditor's recommendation", "auditor recommendation"],
    "audit_observation": ["observation", "observations", "audit observations"],
    "policy_title": ["policy", "policy name"],
    "policy_reference": ["policy content", "policy clause", "policy text"],
    "notable_numbers": ["key figures", "notable figures", "notable number"],
    "criterion": ["criteria"],
    ACTION_PARTY_TIME: [
        "action party & time",
        "action party and time",
        "action party & date",
        "action party/time",
    ],
}
# The header table of a Word audit report (section 8.1).
REPORT_LABELS = {
    "report title": "title",
    "title of report": "title",
    "title": "title",
    "audit period": "audit_period",
    "period": "audit_period",
    "report date": "report_date",
    "date": "report_date",
    "prepared by": "prepared_by",
}
CRITERION_HEADING = re.compile(r"^criteri(?:on|a)\b[\s\d.]*[:\-–—]\s*(?P<title>.+)$", re.IGNORECASE)
DATE_IN_TEXT = re.compile(r"\d{4}-\d{2}-\d{2}|\d{1,2}[/-]\d{1,2}[/-]\d{4}|\d{1,2} [A-Za-z]{3,9} \d{4}")


class ImportFileError(Exception):
    """The file as a whole cannot be read. Nothing is imported."""


@dataclass
class ParsedRow:
    line: int
    data: dict
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _normal(label: str) -> str:
    return re.sub(r"\s+", " ", label.replace("_", " ").strip().strip(":").lower())


def _header_aliases() -> dict[str, set[str]]:
    custom = getattr(settings, "IMPORT_HEADER_ALIASES", {}) or {}
    aliases = {}
    for f in (*FIELDS, *OPTIONAL_FIELDS, ACTION_PARTY_TIME):
        names = {f, *BUILTIN_ALIASES.get(f, []), *custom.get(f, [])}
        aliases[f] = {_normal(n) for n in names}
    return aliases


def _field_for(label: str, aliases: dict[str, set[str]]) -> str | None:
    label = _normal(label)
    for f, names in aliases.items():
        if label in names:
            return f
    return None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(filename: str, content: bytes) -> list[list[str]]:
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]
    if name.endswith(".xlsx"):
        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for damaged files
            raise ImportFileError(
                "This Excel file could not be opened. Save it again as .xlsx and retry."
            ) from exc
        sheet = book.worksheets[0]
        rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        book.close()
        return rows
    raise ImportFileError("Upload the final queries as a Word (.docx), Excel (.xlsx) or CSV file.")


def _parse_date(value: str) -> date | None:
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    return None


def split_action_party_time(value: str) -> tuple[str, str]:
    """'Finance Manager, 31/01/2027' -> ('Finance Manager', '31/01/2027'). No date found: all of it is the
    party."""
    match = DATE_IN_TEXT.search(value)
    if not match:
        return value.strip(), ""
    party = (value[: match.start()] + value[match.end() :]).strip(" ,;-–—\t\n").strip()
    party = re.sub(r"\s*\b(by|before|on|due)\s*$", "", party, flags=re.IGNORECASE).strip(" ,;-–—")
    return party, match.group(0)


def _complete(data: dict) -> dict:
    """Every field present, and a combined action-party-and-time value split into its two fields."""
    combined = data.pop(ACTION_PARTY_TIME, "")
    if combined:
        party, when = split_action_party_time(combined)
        if not data.get("responsible_officer"):
            data["responsible_officer"] = party
        if not data.get("deadline"):
            data["deadline"] = when
    return {f: (data.get(f) or "").strip() for f in (*FIELDS, *OPTIONAL_FIELDS)}


def parse_rows(table: list[list[str]], *, first_line: int = 1, criterion: str = "") -> list[ParsedRow]:
    rows = [r for r in table if any(c.strip() for c in r)]
    if not rows:
        raise ImportFileError("The file is empty.")
    aliases = _header_aliases()
    index = {}
    for i, h in enumerate(rows[0]):
        f = _field_for(h, aliases)
        if f and f not in index:
            index[f] = i
    covered = set(index)
    if ACTION_PARTY_TIME in index:
        covered |= {"responsible_officer", "deadline"}
    missing = [f for f in FIELDS if f not in covered]
    if missing:
        raise ImportFileError(f"The file is missing these columns: {', '.join(missing)}.")
    body = rows[1:]
    if len(body) > settings.IMPORT_MAX_ROWS:
        raise ImportFileError(
            f"The file has {len(body)} rows. The limit is {settings.IMPORT_MAX_ROWS} per import."
        )
    line_of = {id(r): i + first_line for i, r in enumerate(table)}
    parsed = []
    for r in body:
        data = {f: (r[i].strip() if i < len(r) else "") for f, i in index.items()}
        if criterion and not data.get("criterion"):
            data["criterion"] = criterion
        parsed.append(ParsedRow(line=line_of[id(r)], data=_complete(data)))
    return parsed


# Word -------------------------------------------------------------------------------------------------


def _blocks(document):
    """Paragraphs and tables in the order they appear in the document body."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def _table_cells(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            # A merged cell is returned once per grid column it spans.
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text.strip())
        rows.append(cells)
    return rows


def _title_from(text: str) -> str:
    first = re.split(r"(?<=[.!?])\s", text.strip(), maxsplit=1)[0]
    return (first[:117] + "...") if len(first) > 120 else first


def read_word(content: bytes) -> tuple[list[ParsedRow], dict]:
    """Queries and report details from a Word audit report. Returns (rows, report) where report holds the
    header values found and the criteria in order, each with its description."""
    try:
        import docx

        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # python-docx raises many types for damaged or non-Word files
        raise ImportFileError(
            "This Word file could not be opened. Save it again as .docx and retry."
        ) from exc

    aliases = _header_aliases()
    report: dict = {"criteria": []}
    criterion: dict | None = None
    describing = False
    rows: list[ParsedRow] = []
    styled_title = ""

    for block in _blocks(document):
        if block.__class__.__name__ == "Paragraph":
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "").lower()
            match = CRITERION_HEADING.match(text)
            if match:
                criterion = {"title": match.group("title").strip()[:300], "description": ""}
                report["criteria"].append(criterion)
                describing = True
            elif describing and criterion is not None and not style.startswith("heading"):
                criterion["description"] = (criterion["description"] + "\n" + text).strip()
            elif style == "title" and not styled_title:
                styled_title = text[:300]
            else:
                describing = False
            continue

        describing = False
        cells = [r for r in _table_cells(block) if any(c for c in r)]
        if not cells:
            continue
        header_fields = {_field_for(c, aliases) for c in cells[0]}
        if len(cells[0]) > 2 and "reference" in header_fields:
            # One table, a query per row.
            more = parse_rows(
                cells, first_line=len(rows) + 1, criterion=criterion["title"] if criterion else ""
            )
            for i, row in enumerate(more):
                row.line = len(rows) + i + 1
            rows.extend(more)
            continue
        pairs = [(r[0], " ".join(c for c in r[1:] if c).strip()) for r in cells if len(r) >= 2]
        fields = {}
        for label, value in pairs:
            f = _field_for(label, aliases)
            if f and f not in fields:
                fields[f] = value
        if "reference" in fields:
            if criterion and not fields.get("criterion"):
                fields["criterion"] = criterion["title"]
            row = ParsedRow(line=len(rows) + 1, data=_complete(fields))
            if not row.data["title"] and row.data["description"]:
                row.data["title"] = _title_from(row.data["description"])
                row.warnings.append("No title given: the first sentence of the query is used")
            rows.append(row)
            continue
        for label, value in pairs:
            key = REPORT_LABELS.get(_normal(label))
            if key and value and not report.get(key):
                report[key] = value[:300]

    # A "Report title" in the header table wins over the document's title paragraph.
    if not report.get("title") and styled_title:
        report["title"] = styled_title
    if not rows:
        raise ImportFileError(
            "No queries were found in this Word file. Put each query in a table with a Reference row, "
            "or use one table with a header row. See the import template."
        )
    if len(rows) > settings.IMPORT_MAX_ROWS:
        raise ImportFileError(f"The file has {len(rows)} queries. The limit is {settings.IMPORT_MAX_ROWS}.")
    return rows, report


def validate(rows: list[ParsedRow], existing_references: set[str], today: date) -> list[ParsedRow]:
    """Blocking errors stop the import. Warnings are shown but do not block (design.md 13)."""
    seen: set[str] = set()
    for row in rows:
        d = row.data
        ref = d["reference"]
        if not ref:
            row.errors.append("Missing reference")
        elif ref in existing_references:
            row.errors.append("Reference already imported in an earlier batch")
        elif ref in seen:
            row.errors.append("Duplicate reference in this file")
        seen.add(ref)

        for f, label in (
            ("title", "title"),
            ("description", "description"),
            ("responsible_officer", "responsible officer"),
            ("department", "department"),
        ):
            if not d[f]:
                row.errors.append(f"Missing {label}")
        for f, limit in MAX_LENGTH.items():
            if len(d[f]) > limit:
                row.errors.append(f"{f.replace('_', ' ').capitalize()} is over {limit} characters")

        risk = RISKS.get(d["risk_rating"].strip().lower())
        if risk:
            d["risk_rating"] = risk
        else:
            row.errors.append(
                f'Unknown risk rating "{d["risk_rating"] or "blank"}". Use Low, Medium, High or Critical'
            )

        deadline = _parse_date(d["deadline"])
        if deadline is None:
            row.errors.append(f'Bad date "{d["deadline"] or "blank"}". Use YYYY-MM-DD or DD/MM/YYYY')
        else:
            d["deadline"] = deadline.isoformat()
            if deadline < today:
                row.warnings.append("Deadline is before the import date")
    return rows


def parse_report_date(value: str) -> date | None:
    return _parse_date(value) if value else None
